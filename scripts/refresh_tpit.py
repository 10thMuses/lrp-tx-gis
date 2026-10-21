#!/usr/bin/env python3
"""Attribute refresh of the `tpit_lines` layer from the latest ERCOT TPIT.

SCOPE / LIMITATION (read first)
    The existing tpit_subs / tpit_lines geometry was NOT produced by any script
    in this repo. Reverse-engineering shows: every tpit_subs row is an OSM
    `substations` row (identical name/lat/lon/operator/voltage), and every
    tpit_lines feature is a straight 2-point line between two such substations
    (`from` / `to` = the matched OSM substation names), selected from the ~1,340
    TPIT projects that carry both terminals. The terminal->substation matching
    rule cannot be reproduced: the repo's own fuzzy kernel (WRatio>=88,
    same-county or statewide) agrees with only ~51-66% of the existing 266
    endpoints, so deriving NEW lines/substations would require hand curation.
    This script therefore does NOT create geometry. It keeps the existing
    geometry + matched `from`/`to` (keyed on ERCOT project number `pid`) and
    refreshes the TPIT attributes (status, tier, sheet, in_service, owner,
    voltage, county, name, from_raw/to_raw) from the current report.
    tpit_subs carries no TPIT attributes (pure OSM copies) -> nothing to refresh.

Source
    ERCOT Planning page https://www.ercot.com/gridinfo/planning, link titled
    "Transmission Project and Information Tracking" (the xlsx is replaced in
    place; its Last-Modified header dates the report). Sheets parsed:
    Future*/Planned*/Completed*/Cancelled* (header = row 2).

Output (outputs/refresh/)
    tpit_lines_<date>.geojson   merge-ready: python3 build.py merge tpit_lines <file>
    tpit_lines_<date>.meta.json
Flags: --dry-run, --xlsx PATH, --date YYYY-MM-DD
Cancelled projects are dropped from the output (listed by pid in the meta).
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
import urllib.request
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
GEOMS = REPO / 'combined_geoms.geojson'
OUT_DIR = REPO / 'outputs' / 'refresh'
CACHE_DIR = REPO / 'tmp_tpit_cache'  # matches .gitignore `tmp*/`
PLANNING_URL = 'https://www.ercot.com/gridinfo/planning'
FALLBACK_URL = ('https://www.ercot.com/files/docs/2022/03/02/'
                'ERCOT-July-Ad-Hoc-TPIT-No-Cost-071326-UPDATE.xlsx')
LAYER = 'tpit_lines'
UA = {'User-Agent': 'Mozilla/5.0'}
SHEET_PRIORITY = ['Planned', 'Future', 'Completed', 'Cancelled']


def http(url, attempts=4, timeout=180):
    last = None
    for i in range(attempts):
        try:
            req = urllib.request.Request(url, headers=UA)
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return r.read(), dict(r.headers)
        except Exception as e:  # noqa: BLE001
            last = e
            time.sleep(5 * (i + 1))
    raise RuntimeError(f'GET failed {url}: {last}')


def find_tpit_url() -> str:
    try:
        html = http(PLANNING_URL)[0].decode('utf8', 'ignore')
        for m in re.finditer(
                r'<a [^>]*href="([^"]+\.xlsx)"[^>]*title="([^"]*)"', html):
            url, title = m.groups()
            if title.strip() == 'Transmission Project and Information Tracking':
                return url
    except Exception as e:  # noqa: BLE001
        print(f'  planning page lookup failed ({e}); using fallback URL',
              file=sys.stderr)
    return FALLBACK_URL


def download() -> tuple[Path, str, str]:
    url = find_tpit_url()
    blob, hdr = http(url)
    if blob[:2] != b'PK':
        raise RuntimeError('TPIT download is not an xlsx')
    CACHE_DIR.mkdir(exist_ok=True)
    dest = CACHE_DIR / Path(url).name
    tmp = dest.with_suffix('.tmp')
    tmp.write_bytes(blob)
    os.replace(tmp, dest)
    lm = {k.lower(): v for k, v in hdr.items()}.get('last-modified', '')
    return dest, url, lm


def clean(v) -> str:
    if v is None or (isinstance(v, float) and v != v):
        return ''
    return re.sub(r'_x000D_', '', str(v)).strip()


def fmt_kv(v) -> str:
    try:
        f = float(v)
    except (TypeError, ValueError):
        return ''
    return str(int(f)) if f == int(f) else str(f)


def fmt_date(v) -> str:
    return str(v) if isinstance(v, datetime) else ''


def parse_tpit(xlsx: Path) -> dict[str, dict]:
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    by_pid: dict[str, dict] = {}
    for tag in SHEET_PRIORITY:
        sn = next((n for n in wb.sheetnames if n.startswith(tag + 'TPIT')),
                  None)
        if sn is None:
            raise RuntimeError(f'sheet {tag}TPIT* missing: {wb.sheetnames}')
        rows = list(wb[sn].iter_rows(values_only=True))
        hdr = [re.sub(r'\s+', ' ', clean(h)) for h in rows[1]]

        def col(prefix):
            return next(i for i, h in enumerate(hdr) if h.startswith(prefix))
        ix = dict(
            pid=col('ERCOT Project Number'), title=col('Project Title'),
            desc=col('Project Description'), frm=col('Terminal "from"'),
            to=col('Terminal "to"'), status=col('Transmission Status'),
            owner=col('Transmission Owner ('),
            proj=col('Projected In-Service'), act=col('Actual In-Service'),
            kv=col('Service Level kV'), county=col('County Location for Sub'),
            tier=col('Planning Charter Tier'))
        for r in rows[2:]:
            pid = clean(r[ix['pid']])
            if not pid or pid in by_pid:  # first (highest-priority) sheet wins
                continue
            by_pid[pid] = dict(
                sheet=tag, name=clean(r[ix['desc']]) or clean(r[ix['title']]),
                from_raw=clean(r[ix['frm']]), to_raw=clean(r[ix['to']]),
                status=clean(r[ix['status']]), owner=clean(r[ix['owner']]),
                in_service=fmt_date(r[ix['proj']]) or fmt_date(r[ix['act']]),
                voltage=fmt_kv(r[ix['kv']]), county=clean(r[ix['county']]),
                tier=clean(r[ix['tier']]),
                _both=bool(clean(r[ix['frm']]) and clean(r[ix['to']])))
    return by_pid


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('--dry-run', action='store_true')
    ap.add_argument('--xlsx', type=Path)
    ap.add_argument('--date', default=date.today().isoformat())
    a = ap.parse_args()

    if a.xlsx:
        xlsx, url, lm = a.xlsx, str(a.xlsx), ''
    else:
        xlsx, url, lm = download()
    print(f'source : {url}\nlast-modified: {lm or "n/a"}')
    tp = parse_tpit(xlsx)
    print('TPIT projects by sheet:', dict(Counter(v['sheet']
                                                  for v in tp.values())))

    with open(GEOMS, encoding='utf-8') as f:
        feats = json.load(f)['features']
    lines = [f for f in feats if f['properties'].get('layer_id') == LAYER]
    print(f'existing {LAYER}: {len(lines)} features')

    out, dropped_cancel, missing = [], [], []
    changed = Counter()
    keys = ['name', 'owner', 'county', 'from_raw', 'to_raw', 'voltage',
            'status', 'tier', 'sheet', 'in_service']
    for ft in lines:
        p = ft['properties']
        t = tp.get(str(p['pid']))
        if t is None:
            missing.append(p['pid'])
            out.append(ft)  # keep as-is; flagged
            continue
        if t['sheet'] == 'Cancelled':
            dropped_cancel.append(p['pid'])
            continue
        newp = dict(p)
        for k in keys:
            v = t[k]
            if k in ('name', 'owner', 'county', 'tier') and not v:
                v = p.get(k, '')
            if str(p.get(k, '')) != v:
                changed[k] += 1
            newp[k] = v
        out.append({'type': 'Feature', 'properties': newp,
                    'geometry': ft['geometry']})

    both_total = sum(1 for v in tp.values()
                     if v['_both'] and v['sheet'] != 'Cancelled')
    new_pids = [k for k, v in tp.items() if v['_both']
                and v['sheet'] != 'Cancelled'
                and k not in {str(f['properties']['pid']) for f in lines}]
    print(f'features out: {len(out)}  (cancelled dropped {len(dropped_cancel)}'
          f', pid not in report {len(missing)})')
    print('attribute changes vs existing:', dict(changed))
    print('status now:', dict(Counter(f['properties']['status']
                                      for f in out)))
    print('sheet now :', dict(Counter(f['properties']['sheet']
                                      for f in out)))
    print(f'NOT addable without substation matching: {len(new_pids)} further '
          f'non-cancelled TPIT projects have both terminals '
          f'({both_total} total with both terminals vs {len(lines)} in layer)')

    if a.dry_run:
        print('--dry-run: nothing written')
        return 0
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    dest = OUT_DIR / f'tpit_lines_{a.date}.geojson'
    tmp = dest.with_suffix('.geojson.tmp')
    tmp.write_text(json.dumps({'type': 'FeatureCollection', 'features': out},
                              separators=(',', ':')), encoding='utf-8')
    os.replace(tmp, dest)
    meta = dict(layer_id=LAYER, features=len(out), source_url=url,
                source_last_modified=lm, run_date=a.date,
                attribute_changes=dict(changed),
                dropped_cancelled_pids=dropped_cancel,
                pids_missing_from_report=missing,
                geometry='carried forward from combined_geoms.geojson by pid')
    mt = dest.with_suffix('.meta.json.tmp')
    mt.write_text(json.dumps(meta, indent=2))
    os.replace(mt, dest.with_suffix('.meta.json'))
    print(f'wrote {dest}\nmerge: python3 build.py merge {LAYER} '
          f'outputs/refresh/tpit_lines_{a.date}.geojson')
    return 0


if __name__ == '__main__':
    sys.exit(main())
