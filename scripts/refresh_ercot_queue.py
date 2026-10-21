#!/usr/bin/env python3
"""Refresh the `ercot_queue` layer from the latest ERCOT GIS Report.

Fetch -> parse -> map to combined_points.csv schema -> geocode -> write a
merge-ready CSV. Does NOT touch combined_points.csv (read-only here); merge is
a separate, explicit step:

    python3 scripts/refresh_ercot_queue.py
    python3 build.py merge ercot_queue outputs/refresh/ercot_queue_<date>.csv

Source
    ERCOT MIS report type 15933 ("GIS Report"). Doc list JSON:
      https://www.ercot.com/misapp/servlets/IceDocListJsonWS?reportTypeId=15933
    Download:
      https://www.ercot.com/misdownload/servlets/mirDownload?doclookupId=<DocID>
    The report type also carries "Co-located_Battery_Identification_Report_*";
    only FriendlyName `GIS_Report_*` entries are considered. Latest = max
    PublishDate. Sheet parsed: 'Project Details - Large Gen' (header row is the
    row whose first cell == 'INR'; located dynamically).

Column map (validated against the March-2026 report == existing layer)
    INR->inr  Project Name->name  Interconnecting Entity->entity
    POI Location->poi  County->county  CDR Reporting Zone->zone
    Fuel->fuel  Technology->technology  Capacity (MW)->mw == capacity_mw
    (0 MW -> blank, as in the existing layer; negatives kept).
    Derived:
      under_construction = Yes iff 'Construction Start' populated
      commissioned       = Yes iff 'Approved for Synchronization' populated
                           and dated <= report publish date
      funnel_stage       = 5_commissioned | 4_construction (construction start)
                           | 3_ia_fis_complete (IA Signed + FIS Approved)
                           | 2_ia_fis_pending (IA Signed) | 1_no_ia
      group              = carried forward by INR when known; else best-effort
                           `<FIRST DISTINCTIVE NAME TOKEN>__<COUNTY_UPPER>`
                           (approximation; original keying logic is not in repo)

Scope: statewide; fuels GAS/OTH/SOL/WIN (+NUC by default) -- the fuel set of
the existing layer. --all-fuels keeps HYD/OIL/WAT/MWH etc.

Coordinates (same method as scripts/geocode_ercot_queue.py)
    1. INR already in combined_points.csv with same county -> carry lat/lon and
       coords_source forward unchanged.
    2. New INR (or county changed) -> Stage 1 EIA-860 / USWTDB name+county
       fuzzy match, Stage 2 TPIT-POI / OSM-substation-POI match, dc_anchors
       alias match (all imported from geocode_ercot_queue.py, WRatio >= 88).
    3. Otherwise TIGER-2024 county polygon centroid + deterministic +/-0.02 deg
       jitter keyed on INR; coords_source = county_centroid. (The original
       centroid table/jitter is not in the repo; centroid+jitter reproduces its
       statistics, not its exact values.)
    Rows whose county cannot be resolved are dropped and listed in the meta.

Outputs (outputs/refresh/)
    ercot_queue_<date>.csv        merge-ready, combined_points.csv header
    ercot_queue_<date>.meta.json  source report name/date/URL/DocID, counts
<date> is the run date (fetch date, repo convention); the ERCOT report
month/publish date is in the meta file and the printed summary.

Flags: --dry-run (no files written), --xlsx PATH (use a local report; pair with
--report-name/--publish-date), --date YYYY-MM-DD, --all-fuels, --no-compare.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import sys
import time
import urllib.request
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / 'scripts'))
import geocode_ercot_queue as G  # noqa: E402  (stage matchers + TIGER helpers)

COMBINED = REPO / 'combined_points.csv'
OUT_DIR = REPO / 'outputs' / 'refresh'
CACHE_DIR = REPO / 'tmp_ercot_cache'  # matches .gitignore `tmp*/`
LAYER = 'ercot_queue'
SHEET = 'Project Details - Large Gen'
DOCLIST_URL = ('https://www.ercot.com/misapp/servlets/'
               'IceDocListJsonWS?reportTypeId=15933')
DOWNLOAD_URL = ('https://www.ercot.com/misdownload/servlets/'
                'mirDownload?doclookupId={}')
UA = {'User-Agent': 'Mozilla/5.0'}

BASE_FUELS = {'GAS', 'OTH', 'SOL', 'WIN'}
DEFAULT_FUELS = BASE_FUELS | {'NUC'}
REQUIRED_COLS = [
    'INR', 'Project Name', 'GIM Study Phase', 'Interconnecting Entity',
    'POI Location', 'County', 'CDR Reporting Zone', 'Fuel', 'Technology',
    'Capacity (MW)', 'FIS Approved', 'IA Signed', 'Construction Start',
    'Approved for Synchronization',
]
PECOS = 'PECOS'
ADJACENT = ['REEVES', 'CRANE', 'WARD', 'UPTON', 'ECTOR', 'CROCKETT']
JITTER_DEG = 0.02

GROUP_STOP = {
    'SOLAR', 'STORAGE', 'BESS', 'ENERGY', 'WIND', 'POWER', 'PROJECT', 'THE',
    'BATTERY', 'GAS', 'PHASE', 'LLC', 'FARM', 'PLANT', 'STATION', 'FACILITY',
    'GENERATION', 'GENERATING', 'AND', 'OF', 'I', 'II', 'III', 'IV', 'V',
}


# ---------------------------------------------------------------- fetch ----

def http_get(url: str, attempts: int = 4, timeout: int = 180) -> bytes:
    last = None
    for i in range(attempts):
        try:
            req = urllib.request.Request(url, headers=UA)
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return r.read()
        except Exception as e:  # noqa: BLE001
            last = e
            time.sleep(5 * (i + 1))
    raise RuntimeError(f'GET failed {url}: {last}')


def latest_gis_report() -> dict:
    j = json.loads(http_get(DOCLIST_URL))
    docs = [d['Document'] for d in j['ListDocsByRptTypeRes']['DocumentList']]
    docs = [d for d in docs
            if (d.get('FriendlyName') or '').startswith('GIS_Report')
            and d.get('Extension') == 'xlsx']
    if not docs:
        raise RuntimeError('no GIS_Report_* documents in ERCOT doc list')
    docs.sort(key=lambda d: d['PublishDate'], reverse=True)
    return docs[0]


def download_report(doc: dict) -> Path:
    CACHE_DIR.mkdir(exist_ok=True)
    dest = CACHE_DIR / f"{doc['FriendlyName']}_{doc['DocID']}.xlsx"
    if dest.exists() and dest.stat().st_size > 100_000:
        return dest
    for _ in range(3):
        blob = http_get(DOWNLOAD_URL.format(doc['DocID']))
        if blob[:2] == b'PK':  # xlsx is a zip; ERCOT sometimes returns HTML
            tmp = dest.with_suffix('.tmp')
            tmp.write_bytes(blob)
            os.replace(tmp, dest)
            return dest
        time.sleep(10)
    raise RuntimeError(f'download did not return an xlsx for {doc["DocID"]}')


# ---------------------------------------------------------------- parse ----

def s(v) -> str:
    return '' if v is None else str(v).strip()


def to_dt(v):
    if isinstance(v, datetime):
        return v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    return None


def parse_large_gen(xlsx: Path) -> list[dict]:
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise RuntimeError(f'sheet {SHEET!r} missing; have {wb.sheetnames}')
    rows = list(wb[SHEET].iter_rows(values_only=True))
    hdr_i = next((i for i, r in enumerate(rows[:80]) if s(r[0]) == 'INR'), None)
    if hdr_i is None:
        raise RuntimeError("header row (first cell 'INR') not found")
    hdr = [re.sub(r'\s+', ' ', s(h)) for h in rows[hdr_i]]
    missing = [c for c in REQUIRED_COLS if c not in hdr]
    if missing:
        raise RuntimeError(f'ERCOT column rename? missing: {missing}')
    pos = {h: i for i, h in enumerate(hdr) if h}
    out = []
    for r in rows[hdr_i + 1:]:
        inr, nm = s(r[pos['INR']]), s(r[pos['Project Name']])
        if not inr or not nm:
            continue
        rec = {c: r[pos[c]] for c in REQUIRED_COLS}
        out.append(rec)
    return out


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def fmt_num(x):
    return '' if x is None else format(round(x, 4), '.10g')


def derive_group(name: str, entity: str, county: str) -> str:
    toks = re.sub(r'[^A-Z0-9 ]', ' ', name.upper()).split()
    pick = next((t for t in toks if t.isalpha() and len(t) >= 3
                 and t not in GROUP_STOP), None)
    if pick is None:
        toks = re.sub(r'[^A-Z0-9 ]', ' ', entity.upper()).split()
        pick = next((t for t in toks if t.isalpha() and len(t) >= 3
                     and t not in GROUP_STOP), 'UNKNOWN')
    return f"{pick}__{county.upper().replace(' ', '_')}"


def build_row(rec: dict, publish: datetime) -> dict:
    mw = num(rec['Capacity (MW)'])
    mw = None if (mw is None or mw == 0) else mw
    cstart = to_dt(rec['Construction Start'])
    sync = to_dt(rec['Approved for Synchronization'])
    ia = to_dt(rec['IA Signed'])
    fis = to_dt(rec['FIS Approved'])
    commissioned = bool(sync and sync <= publish)
    if commissioned:
        stage = '5_commissioned'
    elif cstart:
        stage = '4_construction'
    elif ia and fis:
        stage = '3_ia_fis_complete'
    elif ia:
        stage = '2_ia_fis_pending'
    else:
        stage = '1_no_ia'
    return {
        'name': s(rec['Project Name']), 'county': s(rec['County']),
        'technology': s(rec['Technology']), 'inr': s(rec['INR']),
        'fuel': s(rec['Fuel']), 'mw': fmt_num(mw), 'capacity_mw': fmt_num(mw),
        'zone': s(rec['CDR Reporting Zone']), 'poi': s(rec['POI Location']),
        'entity': s(rec['Interconnecting Entity']), 'funnel_stage': stage,
        'under_construction': 'Yes' if cstart else 'No',
        'commissioned': 'Yes' if commissioned else 'No',
    }


# ------------------------------------------------------------- geocode ----

def load_existing() -> dict:
    ex = {}
    with open(COMBINED, newline='', encoding='utf-8') as f:
        for row in csv.DictReader(f):
            if row.get('layer_id') == LAYER:
                ex[row['inr'].strip()] = row
    return ex


def county_key(c: str) -> str:
    return re.sub(r"[^A-Z0-9]", '', c.upper())


def jitter(inr: str):
    h = hashlib.sha1(inr.encode()).digest()
    u = int.from_bytes(h[:4], 'big') / 2**32
    v = int.from_bytes(h[4:8], 'big') / 2**32
    return (u * 2 - 1) * JITTER_DEG, (v * 2 - 1) * JITTER_DEG


class Geocoder:
    def __init__(self):
        self.eia, *_ = G.build_eia_index()
        self.usw, _ = G.build_uswtdb_index()
        self.tree, self.polys, self.cnames = G.build_tx_county_polys()
        self.tpit, *_ = G.build_substation_index_from_layer(
            'tpit_subs', self.tree, self.polys, self.cnames)
        self.osm, *_ = G.build_substation_index_from_layer(
            'substations', self.tree, self.polys, self.cnames)
        self.dc, _ = G.build_dc_anchors_index()
        self.centroid = {}
        for poly, nm in zip(self.polys, self.cnames):
            c = poly.centroid
            self.centroid[county_key(nm)] = (c.y, c.x)

    def locate(self, row: dict):
        """Return (lat, lon, coords_source) or None if county unresolved."""
        co = G.norm_county(row['county'])
        qn = G.norm_name(row['name'])
        poi_n = G.norm_substation_name(row['poi'])
        bucket = G.fuel_bucket(row['fuel'], row['technology'])
        if qn and co in self.eia:
            m = G.best_match(qn, self.eia[co])
            if m:
                return m[0], m[1], 'eia860'
        if bucket == 'wind' and qn and co in self.usw:
            m = G.best_match(qn, self.usw[co])
            if m:
                return m[0], m[1], 'uswtdb'
        if poi_n and co in self.tpit:
            m = G.best_match(poi_n, self.tpit[co])
            if m:
                return m[0], m[1], 'tpit_poi'
        if poi_n and co in self.osm:
            m = G.best_match(poi_n, self.osm[co])
            if m:
                return m[0], m[1], 'substation_poi'
        if co in self.dc:
            hay = f"{row['entity']} {row['name']}"
            for pat, alat, alon, _ in self.dc[co]:
                if pat.search(hay):
                    return alat, alon, 'dc_anchors'
        c = self.centroid.get(county_key(row['county']))
        if c is None:
            return None
        dy, dx = jitter(row['inr'])
        return c[0] + dy, c[1] + dx, 'county_centroid'


# --------------------------------------------------------------- main ----

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('--dry-run', action='store_true')
    ap.add_argument('--xlsx', type=Path)
    ap.add_argument('--report-name')
    ap.add_argument('--publish-date', help='YYYY-MM-DD (with --xlsx)')
    ap.add_argument('--date', default=date.today().isoformat())
    ap.add_argument('--all-fuels', action='store_true')
    ap.add_argument('--no-compare', action='store_true')
    a = ap.parse_args()

    if a.xlsx:
        xlsx = a.xlsx
        doc = {'FriendlyName': a.report_name or xlsx.stem, 'DocID': '',
               'PublishDate': (a.publish_date or a.date) + 'T00:00:00'}
    else:
        doc = latest_gis_report()
        xlsx = download_report(doc)
    publish = datetime.fromisoformat(doc['PublishDate'][:19])
    print(f"source report : {doc['FriendlyName']}  "
          f"(published {publish.date()}, DocID {doc['DocID'] or 'n/a'})")

    recs = parse_large_gen(xlsx)
    fuels = None if a.all_fuels else DEFAULT_FUELS
    seen, rows, dropped_fuel, dup = set(), [], Counter(), 0
    for rec in recs:
        r = build_row(rec, publish)
        if r['inr'] in seen:
            dup += 1
            continue
        seen.add(r['inr'])
        if fuels is not None and r['fuel'] not in fuels:
            dropped_fuel[r['fuel']] += 1
            continue
        rows.append(r)
    print(f'parsed rows   : {len(recs)}  kept {len(rows)}  '
          f'dup-INR {dup}  fuel-filtered {dict(dropped_fuel)}')

    existing = load_existing()
    print(f'existing layer: {len(existing)} rows')
    geo = Geocoder()
    carried = new_geo = 0
    no_coord = []
    out_rows = []
    for r in rows:
        old = existing.get(r['inr'])
        if (old and G.norm_county(old['county']) == G.norm_county(r['county'])
                and old.get('lat') and old.get('lon')):
            r['lat'], r['lon'] = old['lat'], old['lon']
            r['coords_source'] = old.get('coords_source') or 'county_centroid'
            r['group'] = old.get('group') or derive_group(
                r['name'], r['entity'], r['county'])
            carried += 1
        else:
            loc = geo.locate(r)
            if loc is None:
                no_coord.append(r['inr'])
                continue
            r['lat'], r['lon'] = f'{loc[0]:.6f}', f'{loc[1]:.6f}'
            r['coords_source'] = loc[2]
            r['group'] = derive_group(r['name'], r['entity'], r['county'])
            new_geo += 1
        out_rows.append(r)

    # ---- summary
    def mw(r):
        return float(r['capacity_mw']) if r['capacity_mw'] else 0.0

    print(f'\nrows out      : {len(out_rows)}  (carried-forward coords '
          f'{carried}, newly geocoded {new_geo}, dropped no-county '
          f'{len(no_coord)})')
    print('by funnel_stage:', dict(sorted(Counter(
        r['funnel_stage'] for r in out_rows).items())))
    print('by fuel        :', dict(Counter(r['fuel'] for r in out_rows)))
    cs = Counter(r['coords_source'] for r in out_rows)
    precise = sum(v for k, v in cs.items() if k != 'county_centroid')
    print('coords_source  :', dict(cs))
    print(f'match rate     : {precise}/{len(out_rows)} non-centroid '
          f'({100 * precise / max(1, len(out_rows)):.1f}%)')
    newcs = Counter(r['coords_source'] for r in out_rows
                    if r['inr'] not in existing)
    print('  (new INRs)   :', dict(newcs))
    print(f"total MW       : {sum(mw(r) for r in out_rows):,.0f}  "
          f"(excl. commissioned "
          f"{sum(mw(r) for r in out_rows if r['commissioned'] == 'No'):,.0f})")

    def county_mw(names):
        tot = defaultdict(lambda: [0, 0.0, 0.0])
        for r in out_rows:
            c = r['county'].upper()
            if c in names:
                t = tot[c]
                t[0] += 1
                t[1] += mw(r)
                t[2] += mw(r) if r['commissioned'] == 'No' else 0
        return tot

    for label, names in (('Pecos', {PECOS}), ('Adjacent six', set(ADJACENT))):
        t = county_mw(names)
        print(f'{label:<13s}: n={sum(v[0] for v in t.values())} '
              f'MW={sum(v[1] for v in t.values()):,.0f} '
              f'(excl. commissioned {sum(v[2] for v in t.values()):,.0f})')
        if label != 'Pecos':
            for c in ADJACENT:
                v = t.get(c, [0, 0, 0])
                print(f'    {c:<9s} n={v[0]:<4d} MW={v[1]:>9,.0f}')

    if not a.no_compare:
        cur = {r['inr'] for r in out_rows}
        old_set = set(existing)
        print(f'\nvs existing    : +{len(cur - old_set)} new INR, '
              f'-{len(old_set - cur)} gone, {len(cur & old_set)} retained; '
              f'rows {len(existing)} -> {len(out_rows)} '
              f'({len(out_rows) - len(existing):+d}); MW '
              f'{sum(float(v["capacity_mw"] or 0) for v in existing.values()):,.0f}'
              f' -> {sum(mw(r) for r in out_rows):,.0f}')
        agree = Counter()
        for r in out_rows:
            o = existing.get(r['inr'])
            if not o:
                continue
            for k in ('funnel_stage', 'commissioned', 'under_construction',
                      'group', 'poi', 'entity', 'name', 'fuel'):
                agree[k] += (s(o.get(k)) == r[k])
        n = len(cur & old_set) or 1
        print('field agreement on retained INRs (existing vs refreshed): ' +
              ', '.join(f'{k} {100 * v / n:.0f}%' for k, v in agree.items()))
    if no_coord:
        print(f'dropped (county unresolved) INRs: {no_coord}')

    if a.dry_run:
        print('\n--dry-run: nothing written')
        return 0

    with open(COMBINED, newline='', encoding='utf-8') as f:
        header = next(csv.reader(f))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / f'ercot_queue_{a.date}.csv'
    tmp = out.with_suffix('.csv.tmp')
    with open(tmp, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=header, extrasaction='ignore',
                           lineterminator='\n')
        w.writeheader()
        for r in out_rows:
            r['layer_id'] = LAYER
            w.writerow(r)
    os.replace(tmp, out)
    meta = {
        'layer_id': LAYER, 'rows': len(out_rows),
        'source_report': doc['FriendlyName'], 'source_doc_id': doc['DocID'],
        'source_publish_date': str(publish.date()),
        'source_list_url': DOCLIST_URL,
        'source_download_url': DOWNLOAD_URL.format(doc['DocID']),
        'sheet': SHEET, 'run_date': a.date,
        'fuels': 'all' if a.all_fuels else sorted(fuels),
        'coords_source_counts': dict(cs), 'carried_forward': carried,
        'newly_geocoded': new_geo, 'dropped_no_county_inrs': no_coord,
    }
    mtmp = out.with_suffix('.meta.json.tmp')
    mtmp.write_text(json.dumps(meta, indent=2))
    os.replace(mtmp, out.with_suffix('.meta.json'))
    print(f'\nwrote {out}  ({out.stat().st_size / 1e3:.0f} KB)')
    print(f'merge: python3 build.py merge {LAYER} '
          f'outputs/refresh/ercot_queue_{a.date}.csv')
    return 0


if __name__ == '__main__':
    sys.exit(main())
